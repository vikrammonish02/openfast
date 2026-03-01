"""IEC 61400-1 compliant loads post-processing service.

Processes real OpenFAST simulation output files (.out/.outb) for:
  - Extreme loads table (max/min per channel with concurrent loads, safety factors)
  - Fatigue DEL table (per channel per Wöhler exponent with weighted combination)
  - Statistics summary (per-channel aggregate across all cases)
  - Excel export for Components department

Reuses existing validated modules:
  - ExtremeLoadExtractor  (apps/api/app/postprocessing/extreme_loads.py)
  - DELCalculator         (apps/api/app/postprocessing/del_calculator.py)
  - StatisticsCalculator  (apps/api/app/postprocessing/statistics.py)
  - OutputReader          (apps/api/app/openfast/output_reader.py)

References
----------
IEC 61400-1:2019  Sections 7.6 (Ultimate), 7.6.3 (Fatigue)
"""

from __future__ import annotations

import io
import logging
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

import numpy as np

logger = logging.getLogger("windforge.iec_loads")


# ---------------------------------------------------------------------------
# Data classes
# ---------------------------------------------------------------------------
@dataclass
class CaseConfig:
    """Configuration for a single simulation case to include in analysis."""

    case_id: str
    output_path: str
    dlc_number: str
    wind_speed: float
    seed: int
    yaw_misalignment: float
    safety_factor: float
    probability_weight: float
    analysis_type: str  # "fatigue" or "ultimate"


@dataclass
class ExtremeLoadRow:
    """A single row of the IEC extreme loads table."""

    channel: str
    unit: str
    max_characteristic: float
    max_design: float
    max_dlc: str
    max_vhub: float
    max_time: float
    max_case_id: str
    min_characteristic: float
    min_design: float
    min_dlc: str
    min_vhub: float
    min_time: float
    min_case_id: str
    safety_factor_max: float
    safety_factor_min: float


@dataclass
class ConcurrentLoadEntry:
    """Concurrent channel values at the timestep of a governing extreme."""

    governing_channel: str
    extreme_type: str  # "max" or "min"
    timestep_values: dict[str, float]


@dataclass
class DELRow:
    """A single row of the fatigue DEL table."""

    channel: str
    unit: str
    del_values: dict[str, float]  # "m=3" -> value
    n_equivalent: float


@dataclass
class StatisticsRow:
    """A single row of the statistics summary table."""

    channel: str
    unit: str
    mean: float
    std: float
    min_val: float
    max_val: float
    abs_max: float
    n_cases: int


@dataclass
class CaseSummaryRow:
    """Summary of a single case included in the analysis."""

    case_id: str
    dlc_number: str
    wind_speed: float
    seed_number: int
    yaw_misalignment: float
    analysis_type: str
    safety_factor: float
    probability_weight: float


@dataclass
class IECLoadsResult:
    """Complete IEC loads analysis result."""

    simulation_id: str
    simulation_name: str
    n_cases_analyzed: int
    channels_analyzed: list[str]
    extreme_loads: list[ExtremeLoadRow]
    concurrent_loads: list[ConcurrentLoadEntry]
    del_table: list[DELRow]
    statistics_table: list[StatisticsRow]
    case_summary: list[CaseSummaryRow]


# ---------------------------------------------------------------------------
# Main analysis function
# ---------------------------------------------------------------------------
def run_iec_loads_analysis(
    case_configs: list[CaseConfig],
    simulation_id: str = "",
    simulation_name: str = "",
    channels: list[str] | None = None,
    t_start: float = 30.0,
    wohler_exponents: list[float] | None = None,
    n_equivalent: float = 1e7,
    consequence_factor: float = 1.0,
) -> IECLoadsResult:
    """Run full IEC 61400-1 loads analysis on real simulation output files.

    Parameters
    ----------
    case_configs : list[CaseConfig]
        Configuration for each simulation case to process.
    simulation_id : str
        ID of the parent simulation.
    simulation_name : str
        Display name for the simulation.
    channels : list[str] | None
        Specific channels to analyze. None = all common channels.
    t_start : float
        Time (s) to skip initial transient.
    wohler_exponents : list[float] | None
        Wöhler exponents for DEL calculation.
    n_equivalent : float
        Number of equivalent cycles for DEL reference.
    consequence_factor : float
        Consequence of failure factor γn (IEC 61400-1 Table 3).

    Returns
    -------
    IECLoadsResult
        Complete analysis with extreme, fatigue, and statistics tables.
    """
    from app.openfast.output_reader import OutputReader
    from app.postprocessing.del_calculator import DELCalculator
    from app.postprocessing.extreme_loads import (
        ExtremeLoadExtractor,
        SimulationResult,
    )
    from app.postprocessing.statistics import StatisticsCalculator

    if wohler_exponents is None:
        wohler_exponents = [3.0, 4.0, 6.0, 8.0, 10.0, 12.0]

    reader = OutputReader()
    stats_calc = StatisticsCalculator(t_start=t_start)
    del_calc = DELCalculator()
    extreme_extractor = ExtremeLoadExtractor(
        t_start=t_start,
        consequence_factor=consequence_factor,
    )

    # --- Phase 1: Load all output files and determine common channels -------
    loaded_outputs: dict[str, object] = {}  # case_id -> OutputData
    all_channel_sets: list[set[str]] = []
    channel_unit_map: dict[str, str] = {}

    for cfg in case_configs:
        try:
            output_data = reader.load(Path(cfg.output_path))
            loaded_outputs[cfg.case_id] = output_data
            ch_set = set(output_data.channel_names)
            all_channel_sets.append(ch_set)
            for name, unit in zip(output_data.channel_names, output_data.channel_units):
                if name not in channel_unit_map:
                    channel_unit_map[name] = unit
        except Exception as exc:
            logger.warning("Failed to load output for case %s: %s", cfg.case_id, exc)
            continue

    if not loaded_outputs:
        logger.error("No output files could be loaded")
        return IECLoadsResult(
            simulation_id=simulation_id,
            simulation_name=simulation_name,
            n_cases_analyzed=0,
            channels_analyzed=[],
            extreme_loads=[],
            concurrent_loads=[],
            del_table=[],
            statistics_table=[],
            case_summary=[],
        )

    # Determine common channels (intersection across all loaded cases)
    common_channels = set.intersection(*all_channel_sets) if all_channel_sets else set()
    # Remove "Time" from analysis channels
    common_channels.discard("Time")
    common_channels.discard("time")

    if channels is not None:
        # Filter to user-requested channels
        common_channels = common_channels.intersection(set(channels))

    sorted_channels = sorted(common_channels)

    # --- Phase 2: Statistics across all cases --------------------------------
    case_stats_list = []
    for cfg in case_configs:
        if cfg.case_id not in loaded_outputs:
            continue
        output_data = loaded_outputs[cfg.case_id]
        try:
            case_stats = stats_calc.calculate_from_output(
                data=output_data.data,
                channel_names=output_data.channel_names,
                channel_units=output_data.channel_units,
                case_id=cfg.case_id,
            )
            case_stats_list.append(case_stats)
        except Exception as exc:
            logger.warning("Statistics failed for case %s: %s", cfg.case_id, exc)

    aggregated_stats = StatisticsCalculator.aggregate_across_cases(case_stats_list)

    statistics_table: list[StatisticsRow] = []
    for ch_name in sorted_channels:
        if ch_name in aggregated_stats:
            s = aggregated_stats[ch_name]
            statistics_table.append(StatisticsRow(
                channel=ch_name,
                unit=channel_unit_map.get(ch_name, ""),
                mean=s["mean"],
                std=s["std_max"],
                min_val=s["min"],
                max_val=s["max"],
                abs_max=s["abs_max"],
                n_cases=int(s["n_cases"]),
            ))

    # --- Phase 3: Extreme loads across all cases -----------------------------
    sim_results_for_extremes: list[SimulationResult] = []
    for cfg in case_configs:
        if cfg.case_id not in loaded_outputs:
            continue
        output_data = loaded_outputs[cfg.case_id]
        channels_dict = {}
        for idx, name in enumerate(output_data.channel_names):
            if name in common_channels:
                channels_dict[name] = output_data.data[:, idx]

        sim_results_for_extremes.append(SimulationResult(
            case_id=cfg.case_id,
            dlc_number=cfg.dlc_number,
            wind_speed=cfg.wind_speed,
            seed=cfg.seed,
            safety_factor=cfg.safety_factor,
            time=output_data.time,
            data=channels_dict,
            channel_units={
                name: channel_unit_map.get(name, "")
                for name in common_channels
            },
        ))

    extremes_dict = extreme_extractor.extract_extremes(
        sim_results_for_extremes,
        channels=sorted_channels,
    )

    extreme_loads: list[ExtremeLoadRow] = []
    for ch_name in sorted_channels:
        if ch_name not in extremes_dict:
            continue
        ch_ext = extremes_dict[ch_name]
        extreme_loads.append(ExtremeLoadRow(
            channel=ch_name,
            unit=ch_ext.channel_unit,
            max_characteristic=ch_ext.max_extreme.characteristic,
            max_design=ch_ext.max_extreme.design,
            max_dlc=ch_ext.max_extreme.source_dlc,
            max_vhub=ch_ext.max_extreme.wind_speed,
            max_time=ch_ext.max_extreme.time,
            max_case_id=ch_ext.max_extreme.source_case_id,
            min_characteristic=ch_ext.min_extreme.characteristic,
            min_design=ch_ext.min_extreme.design,
            min_dlc=ch_ext.min_extreme.source_dlc,
            min_vhub=ch_ext.min_extreme.wind_speed,
            min_time=ch_ext.min_extreme.time,
            min_case_id=ch_ext.min_extreme.source_case_id,
            safety_factor_max=ch_ext.max_extreme.safety_factor,
            safety_factor_min=ch_ext.min_extreme.safety_factor,
        ))

    # --- Phase 4: Concurrent loads -------------------------------------------
    concurrent_loads: list[ConcurrentLoadEntry] = []
    for ch_name in sorted_channels:
        if ch_name not in extremes_dict:
            continue
        ch_ext = extremes_dict[ch_name]
        for ext_type, extreme in [("max", ch_ext.max_extreme), ("min", ch_ext.min_extreme)]:
            gov_case_id = extreme.source_case_id
            gov_time = extreme.time
            if gov_case_id not in loaded_outputs:
                continue
            gov_output = loaded_outputs[gov_case_id]
            time_arr = gov_output.time
            # Find the closest timestep index
            mask = time_arr >= t_start
            if not np.any(mask):
                mask = np.ones(len(time_arr), dtype=bool)
            filtered_time = time_arr[mask]
            time_idx_filtered = int(np.argmin(np.abs(filtered_time - gov_time)))
            time_idx = int(np.where(mask)[0][time_idx_filtered])

            # Read all channel values at this timestep
            ts_values: dict[str, float] = {}
            for idx, name in enumerate(gov_output.channel_names):
                if name in common_channels and name != ch_name:
                    ts_values[name] = float(gov_output.data[time_idx, idx])

            concurrent_loads.append(ConcurrentLoadEntry(
                governing_channel=ch_name,
                extreme_type=ext_type,
                timestep_values=ts_values,
            ))

    # --- Phase 5: Fatigue DEL (fatigue-type cases only) ----------------------
    fatigue_configs = [c for c in case_configs if c.analysis_type == "fatigue"]
    del_table: list[DELRow] = []

    if fatigue_configs:
        for ch_name in sorted_channels:
            del_vals: dict[str, float] = {}
            for m in wohler_exponents:
                del_per_case: list[float] = []
                case_weights: list[float] = []

                for cfg in fatigue_configs:
                    if cfg.case_id not in loaded_outputs:
                        continue
                    output_data = loaded_outputs[cfg.case_id]

                    # Find the channel index
                    ch_idx = None
                    for idx, name in enumerate(output_data.channel_names):
                        if name == ch_name:
                            ch_idx = idx
                            break
                    if ch_idx is None:
                        continue

                    signal = output_data.data[:, ch_idx]
                    dt = output_data.dt

                    try:
                        del_val = del_calc.calculate_del(
                            signal=signal,
                            dt=dt,
                            m_exponent=m,
                            n_equivalent=n_equivalent,
                            t_start=t_start,
                        )
                        del_per_case.append(del_val)
                        weight = cfg.probability_weight if cfg.probability_weight > 0 else 1.0
                        case_weights.append(weight)
                    except Exception as exc:
                        logger.warning(
                            "DEL failed for %s m=%s case %s: %s",
                            ch_name, m, cfg.case_id, exc,
                        )

                if del_per_case:
                    # Normalize weights
                    total_weight = sum(case_weights)
                    if total_weight > 0:
                        norm_weights = [w / total_weight for w in case_weights]
                    else:
                        norm_weights = [1.0 / len(case_weights)] * len(case_weights)

                    try:
                        combined = DELCalculator.combine_del_across_cases(
                            del_per_case=del_per_case,
                            case_weights=norm_weights,
                            m_exponent=m,
                        )
                        del_vals[f"m={m:.0f}"] = combined
                    except Exception:
                        del_vals[f"m={m:.0f}"] = float("nan")
                else:
                    del_vals[f"m={m:.0f}"] = 0.0

            del_table.append(DELRow(
                channel=ch_name,
                unit=channel_unit_map.get(ch_name, ""),
                del_values=del_vals,
                n_equivalent=n_equivalent,
            ))

    # --- Case summary --------------------------------------------------------
    case_summary: list[CaseSummaryRow] = []
    for cfg in case_configs:
        if cfg.case_id in loaded_outputs:
            case_summary.append(CaseSummaryRow(
                case_id=cfg.case_id,
                dlc_number=cfg.dlc_number,
                wind_speed=cfg.wind_speed,
                seed_number=cfg.seed,
                yaw_misalignment=cfg.yaw_misalignment,
                analysis_type=cfg.analysis_type,
                safety_factor=cfg.safety_factor,
                probability_weight=cfg.probability_weight,
            ))

    return IECLoadsResult(
        simulation_id=simulation_id,
        simulation_name=simulation_name,
        n_cases_analyzed=len(loaded_outputs),
        channels_analyzed=sorted_channels,
        extreme_loads=extreme_loads,
        concurrent_loads=concurrent_loads,
        del_table=del_table,
        statistics_table=statistics_table,
        case_summary=case_summary,
    )


# ---------------------------------------------------------------------------
# Excel export
# ---------------------------------------------------------------------------
def generate_iec_loads_excel(
    result: IECLoadsResult,
    project_name: str = "",
    turbine_info: dict | None = None,
) -> bytes:
    """Generate an IEC-compliant Excel workbook from analysis results.

    Parameters
    ----------
    result : IECLoadsResult
        Analysis results from run_iec_loads_analysis().
    project_name : str
        Project name for the summary sheet.
    turbine_info : dict | None
        Turbine specifications for the summary sheet.

    Returns
    -------
    bytes
        XLSX file contents as bytes.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    header_font = Font(bold=True, size=11)
    header_fill = PatternFill(start_color="2B3A4E", end_color="2B3A4E", fill_type="solid")
    header_font_white = Font(bold=True, size=11, color="FFFFFF")
    title_font = Font(bold=True, size=14)
    subtitle_font = Font(bold=True, size=12)
    num_fmt_4 = "0.0000"
    num_fmt_2 = "0.00"
    num_fmt_1 = "0.0"

    if turbine_info is None:
        turbine_info = {}

    # ===== Sheet 1: Summary =================================================
    ws_sum = wb.active
    ws_sum.title = "Summary"
    ws_sum.sheet_properties.tabColor = "1F4E79"

    ws_sum["A1"] = "IEC 61400-1 Loads Analysis Report"
    ws_sum["A1"].font = title_font
    ws_sum.merge_cells("A1:F1")

    rows = [
        ("Project:", project_name),
        ("Simulation:", result.simulation_name),
        ("Date:", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Cases Analyzed:", result.n_cases_analyzed),
        ("Channels Analyzed:", len(result.channels_analyzed)),
        ("Software:", "WindForge (OpenFAST + openfast_toolbox)"),
        ("Standard:", "IEC 61400-1:2019 Ed.4"),
        ("", ""),
    ]
    for i, (label, value) in enumerate(rows, start=3):
        ws_sum[f"A{i}"] = label
        ws_sum[f"A{i}"].font = Font(bold=True)
        ws_sum[f"B{i}"] = value

    # Turbine info
    row = len(rows) + 4
    ws_sum[f"A{row}"] = "Turbine Specifications"
    ws_sum[f"A{row}"].font = subtitle_font
    row += 1
    turbine_fields = [
        ("Rotor Diameter", turbine_info.get("rotor_diameter", ""), "m"),
        ("Hub Height", turbine_info.get("hub_height", ""), "m"),
        ("Rated Power", turbine_info.get("rated_power", ""), "kW"),
        ("Number of Blades", turbine_info.get("num_blades", ""), ""),
        ("Cut-in Speed", turbine_info.get("cut_in_speed", ""), "m/s"),
        ("Cut-out Speed", turbine_info.get("cut_out_speed", ""), "m/s"),
        ("Wind Class", turbine_info.get("wind_class", ""), ""),
        ("Turbulence Class", turbine_info.get("turbulence_class", ""), ""),
    ]
    for label, value, unit in turbine_fields:
        ws_sum[f"A{row}"] = label
        ws_sum[f"B{row}"] = value
        ws_sum[f"C{row}"] = unit
        row += 1

    ws_sum.column_dimensions["A"].width = 22
    ws_sum.column_dimensions["B"].width = 30
    ws_sum.column_dimensions["C"].width = 10

    # ===== Sheet 2: Extreme Loads ============================================
    ws_ext = wb.create_sheet("Extreme Loads")
    ws_ext.sheet_properties.tabColor = "C0392B"

    ext_headers = [
        "Channel", "Unit",
        "Max Char.", "Max Design", "Max DLC", "Max Vhub [m/s]", "Max Time [s]",
        "Min Char.", "Min Design", "Min DLC", "Min Vhub [m/s]", "Min Time [s]",
        "SF (Max)", "SF (Min)",
    ]
    for col, h in enumerate(ext_headers, 1):
        cell = ws_ext.cell(row=1, column=col, value=h)
        cell.font = header_font_white
        cell.fill = header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    for r, row_data in enumerate(result.extreme_loads, 2):
        values = [
            row_data.channel, row_data.unit,
            row_data.max_characteristic, row_data.max_design, row_data.max_dlc,
            row_data.max_vhub, row_data.max_time,
            row_data.min_characteristic, row_data.min_design, row_data.min_dlc,
            row_data.min_vhub, row_data.min_time,
            row_data.safety_factor_max, row_data.safety_factor_min,
        ]
        for col, val in enumerate(values, 1):
            cell = ws_ext.cell(row=r, column=col, value=val)
            cell.border = thin_border
            if isinstance(val, float):
                cell.number_format = num_fmt_2

    # Auto-size columns
    for col in range(1, len(ext_headers) + 1):
        ws_ext.column_dimensions[get_column_letter(col)].width = max(14, len(ext_headers[col - 1]) + 2)
    ws_ext.freeze_panes = "A2"

    # ===== Concurrent loads sub-section (below extreme table) ================
    if result.concurrent_loads:
        gap_row = len(result.extreme_loads) + 4
        ws_ext.cell(row=gap_row, column=1, value="Concurrent Loads").font = subtitle_font
        gap_row += 1
        ws_ext.cell(row=gap_row, column=1, value="(Values of all channels at the timestep of each governing extreme)").font = Font(italic=True, size=9)
        gap_row += 1

        # Group concurrent by governing channel
        conc_channels = result.channels_analyzed[:50]  # Limit columns for Excel
        conc_headers = ["Governing Channel", "Type"] + conc_channels
        for col, h in enumerate(conc_headers, 1):
            cell = ws_ext.cell(row=gap_row, column=col, value=h)
            cell.font = header_font_white
            cell.fill = PatternFill(start_color="1A5276", end_color="1A5276", fill_type="solid")
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
        gap_row += 1

        for cl in result.concurrent_loads:
            ws_ext.cell(row=gap_row, column=1, value=cl.governing_channel).border = thin_border
            ws_ext.cell(row=gap_row, column=2, value=cl.extreme_type).border = thin_border
            for col_idx, ch in enumerate(conc_channels, 3):
                val = cl.timestep_values.get(ch, "")
                cell = ws_ext.cell(row=gap_row, column=col_idx, value=val)
                cell.border = thin_border
                if isinstance(val, float):
                    cell.number_format = num_fmt_2
            gap_row += 1

    # ===== Sheet 3: Fatigue DEL ==============================================
    ws_del = wb.create_sheet("Fatigue DEL")
    ws_del.sheet_properties.tabColor = "27AE60"

    # Determine Wöhler exponents from first row
    m_keys = []
    if result.del_table:
        m_keys = sorted(result.del_table[0].del_values.keys(), key=lambda k: float(k.split("=")[1]))

    del_headers = ["Channel", "Unit"] + m_keys + ["N_eq"]
    for col, h in enumerate(del_headers, 1):
        cell = ws_del.cell(row=1, column=col, value=h)
        cell.font = header_font_white
        cell.fill = header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center")

    for r, row_data in enumerate(result.del_table, 2):
        ws_del.cell(row=r, column=1, value=row_data.channel).border = thin_border
        ws_del.cell(row=r, column=2, value=row_data.unit).border = thin_border
        for col_idx, mk in enumerate(m_keys, 3):
            val = row_data.del_values.get(mk, 0.0)
            cell = ws_del.cell(row=r, column=col_idx, value=val)
            cell.border = thin_border
            cell.number_format = num_fmt_4
        neq_cell = ws_del.cell(row=r, column=len(m_keys) + 3, value=row_data.n_equivalent)
        neq_cell.border = thin_border
        neq_cell.number_format = "0.0E+00"

    for col in range(1, len(del_headers) + 1):
        ws_del.column_dimensions[get_column_letter(col)].width = max(12, len(del_headers[col - 1]) + 2)
    ws_del.freeze_panes = "A2"

    # ===== Sheet 4: Statistics ===============================================
    ws_stats = wb.create_sheet("Statistics")
    ws_stats.sheet_properties.tabColor = "2980B9"

    stats_headers = ["Channel", "Unit", "Mean", "Std", "Min", "Max", "Abs Max", "N Cases"]
    for col, h in enumerate(stats_headers, 1):
        cell = ws_stats.cell(row=1, column=col, value=h)
        cell.font = header_font_white
        cell.fill = header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center")

    for r, row_data in enumerate(result.statistics_table, 2):
        values = [
            row_data.channel, row_data.unit,
            row_data.mean, row_data.std, row_data.min_val, row_data.max_val,
            row_data.abs_max, row_data.n_cases,
        ]
        for col, val in enumerate(values, 1):
            cell = ws_stats.cell(row=r, column=col, value=val)
            cell.border = thin_border
            if isinstance(val, float):
                cell.number_format = num_fmt_4

    for col in range(1, len(stats_headers) + 1):
        ws_stats.column_dimensions[get_column_letter(col)].width = max(12, len(stats_headers[col - 1]) + 2)
    ws_stats.freeze_panes = "A2"

    # ===== Sheet 5: DLC Cases ================================================
    ws_cases = wb.create_sheet("DLC Cases")
    ws_cases.sheet_properties.tabColor = "8E44AD"

    case_headers = [
        "Case ID", "DLC", "Wind Speed [m/s]", "Seed",
        "Yaw [deg]", "Analysis Type", "Safety Factor", "Probability Weight",
    ]
    for col, h in enumerate(case_headers, 1):
        cell = ws_cases.cell(row=1, column=col, value=h)
        cell.font = header_font_white
        cell.fill = header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center")

    for r, row_data in enumerate(result.case_summary, 2):
        values = [
            row_data.case_id[:8] + "...",  # Shorten UUIDs
            row_data.dlc_number,
            row_data.wind_speed,
            row_data.seed_number,
            row_data.yaw_misalignment,
            row_data.analysis_type,
            row_data.safety_factor,
            row_data.probability_weight,
        ]
        for col, val in enumerate(values, 1):
            cell = ws_cases.cell(row=r, column=col, value=val)
            cell.border = thin_border
            if isinstance(val, float):
                cell.number_format = num_fmt_2

    for col in range(1, len(case_headers) + 1):
        ws_cases.column_dimensions[get_column_letter(col)].width = max(15, len(case_headers[col - 1]) + 2)
    ws_cases.freeze_panes = "A2"

    # --- Write to bytes buffer -----------------------------------------------
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()
